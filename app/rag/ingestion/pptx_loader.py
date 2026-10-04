import hashlib
from pathlib import Path

from pptx import Presentation

from app.rag.models import Document
from app.rag.ingestion.base_loader import (
    BaseDocumentLoader,
)


class PPTXLoader(BaseDocumentLoader):

    supported_extensions = {".pptx"}

    def load(
        self,
        file_path: str,
    ) -> Document:

        path = Path(file_path)

        presentation = Presentation(
            str(path)
        )

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):

            slide_text = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    if shape.text.strip():

                        slide_text.append(
                            shape.text
                        )

            if slide_text:

                slides.append(
                    f"\n[Slide {slide_number}]\n"
                    + "\n".join(slide_text)
                )

        text = "\n".join(slides)

        document_id = hashlib.sha256(
            str(path.resolve()).encode()
        ).hexdigest()[:16]

        return Document(
            document_id=document_id,
            source_path=str(path.resolve()),
            file_name=path.name,
            file_type="pptx",
            text=text,
            metadata={
                "slide_count": len(
                    presentation.slides
                )
            },
        )