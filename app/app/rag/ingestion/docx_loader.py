import hashlib
from pathlib import Path

from docx import Document as DocxDocument

from app.rag.models import Document
from app.rag.ingestion.base_loader import (
    BaseDocumentLoader,
)


class DOCXLoader(BaseDocumentLoader):

    supported_extensions = {".docx"}

    def load(
        self,
        file_path: str,
    ) -> Document:

        path = Path(file_path)

        document = DocxDocument(
            str(path)
        )

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        text = "\n".join(paragraphs)

        document_id = hashlib.sha256(
            str(path.resolve()).encode()
        ).hexdigest()[:16]

        return Document(
            document_id=document_id,
            source_path=str(path.resolve()),
            file_name=path.name,
            file_type="docx",
            text=text,
        )