from __future__ import annotations
from pathlib import Path
import shutil

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "corpus"
CORPUS.mkdir(parents=True, exist_ok=True)

FACTS = {
"phoenix_architecture.txt": """Phoenix AI Evaluation Document

Phoenix AI uses Qdrant as its vector database.
The local generation model is qwen3:4b-instruct served through Ollama.
Dense retrieval uses Nomic Embed Text with 768-dimensional embeddings.
Hybrid retrieval combines dense retrieval and BM25 using Reciprocal Rank Fusion (RRF).
The RAG pipeline supports PDF, DOCX, XLSX, PPTX, CSV and TXT documents.
""",
"phoenix_security.txt": """Phoenix AI Security Policy

Read-only knowledge retrieval is considered a safe action.
Modifying external files or applications requires approval before side effects.
Unsupported application operations must fail closed and must not execute a tool.
Document references take precedence over application keywords when the user is asking a knowledge question about a document.
""",
"phoenix_projects.txt": """Phoenix AI Project Facts

The Planning Agent produces validated structured plans and does not execute tools.
The Memory Agent stores user-scoped memories and supports recall and forgetting.
The Research Agent searches external sources and requires cited synthesis.
The Application Control Agent handles supported application operations subject to security policy.
"""
}

def write_text_files():
    for name, text in FACTS.items():
        (CORPUS/name).write_text(text, encoding="utf-8")

def write_docx():
    from docx import Document
    d=Document()
    d.add_heading("Phoenix AI DOCX Evaluation",1)
    d.add_paragraph("Phoenix AI's document-scoped RAG pipeline supports deterministic document references.")
    d.add_paragraph("The canonical vector database in this evaluation corpus is Qdrant.")
    d.add_paragraph("The generation model is qwen3:4b-instruct through Ollama.")
    d.save(CORPUS/"phoenix_evaluation.docx")

def write_xlsx():
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title="Architecture"
    for row in [["Component","Value"],["Vector Database","Qdrant"],["LLM","qwen3:4b-instruct"],["Embedding Model","nomic-embed-text"],["Embedding Dimension",768],["Dense + Lexical Fusion","RRF"]]: ws.append(row)
    ws2=wb.create_sheet("Sales")
    ws2.append(["Order ID","Region","Product","Units","Revenue"])
    for row in [["S001","West","Engine A",10,12000],["S002","West","Engine B",8,10400],["S003","North","Engine A",12,14400],["S004","South","Engine C",5,7500],["S005","North","Engine B",7,9100]]: ws2.append(row)
    wb.save(CORPUS/"phoenix_evaluation.xlsx")

def write_pptx():
    from pptx import Presentation
    p=Presentation(); s=p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text="Phoenix AI Evaluation"; s.placeholders[1].text="Qdrant\nqwen3:4b-instruct\nHybrid retrieval: Dense + BM25 + RRF"
    s=p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text="Safety"; s.placeholders[1].text="Read-only retrieval is safe.\nExternal modifications require approval."
    p.save(CORPUS/"phoenix_evaluation.pptx")

def write_pdf():
    path=CORPUS/"phoenix_evaluation.pdf"
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        c=canvas.Canvas(str(path), pagesize=A4); t=c.beginText(50,800); t.setFont("Helvetica",12)
        for line in ["Phoenix AI PDF Evaluation","","Vector database: Qdrant","Generation model: qwen3:4b-instruct","Embedding model: nomic-embed-text","Hybrid retrieval: Dense + BM25 + RRF","External modifications require approval."]: t.textLine(line)
        c.drawText(t); c.save()
    except ModuleNotFoundError:
        if path.exists() and path.stat().st_size > 0:
            print("reportlab is not installed; keeping existing PDF fixture.")
        else:
            print("WARNING: reportlab is not installed and no PDF fixture exists. Install it with: pip install reportlab")
            return False
    return True

if __name__=="__main__":
    write_text_files(); write_docx(); write_xlsx(); write_pptx(); write_pdf()
    print(f"Evaluation corpus ready: {CORPUS}")
